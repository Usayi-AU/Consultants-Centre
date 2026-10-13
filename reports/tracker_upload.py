"""Import an uploaded Report Tracker workbook into one quarter of ClientReport.

Columns are located by their header text, so both the original layout
(Client Name, CRM, Operations, ...) and later layouts with extra columns
such as Alternate work.
"""

from dataclasses import dataclass, field
from datetime import date, datetime

from openpyxl import load_workbook

from .models import ClientReport, StatusPhase


HEADER_ALIASES = {
    "client": {"client", "client name", "clients"},
    "crm": {"crm", "main crm"},
    "operations": {"operations", "operations staff", "ops", "operations assignee"},
    "submitted": {"submitted"},
    "reviewed": {"reviewed"},
    "sent": {"sent", "sent to client"},
    "due_date": {"due", "due date"},
    "status": {"status"},
    "notes": {"note", "notes", "comment", "comments"},
}
REQUIRED_COLUMNS = ("client", "crm", "operations")

# Client names used in the workbooks that differ from the names in the tracker.
CLIENT_NAME_ALIASES = {
    "brightmarsh": "Marsh Employees PF",
}

STATUS_TEXT = {
    "complete": StatusPhase.SENT_TO_CLIENT,
    "completed": StatusPhase.SENT_TO_CLIENT,
    "sent": StatusPhase.SENT_TO_CLIENT,
    "sent to client": StatusPhase.SENT_TO_CLIENT,
    "reviewed": StatusPhase.REVIEWED,
    "submitted": StatusPhase.SUBMITTED,
    "pending": StatusPhase.PENDING,
}


class TrackerUploadError(Exception):
    pass


@dataclass
class TrackerUploadResult:
    sheet_name: str = ""
    created: list = field(default_factory=list)
    updated: list = field(default_factory=list)
    removed: list = field(default_factory=list)
    skipped: list = field(default_factory=list)


def name_key(value):
    return "".join(ch for ch in str(value or "").lower() if ch.isalnum())


def header_key(value):
    return " ".join(str(value or "").strip().lower().split())


def is_ticked(value):
    if value is None or value is False:
        return False
    if value is True:
        return True
    return str(value).strip().lower() in {"✓", "✔", "x", "yes", "y", "true", "1", "done"}


def parse_due_date(value, year):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    for fmt in ("%d-%b", "%d %b", "%d-%B", "%d %B"):
        try:
            return datetime.strptime(f"{text} {year}", f"{fmt} %Y").date()
        except ValueError:
            continue
    return None


def find_header(workbook):
    """Return (worksheet, header row number, {column key: index}) for the first sheet with tracker headers."""
    sheets = sorted(workbook.worksheets, key=lambda ws: "report tracker" not in ws.title.lower())
    for worksheet in sheets:
        for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=20, values_only=True), start=1):
            columns = {}
            for index, cell in enumerate(row):
                key = header_key(cell)
                for column, aliases in HEADER_ALIASES.items():
                    if key in aliases and column not in columns:
                        columns[column] = index
            if all(column in columns for column in REQUIRED_COLUMNS):
                return worksheet, row_number, columns
    raise TrackerUploadError(
        "Could not find the tracker headers. The sheet needs at least 'Client Name', 'CRM' and 'Operations' columns."
    )


def status_from_row(row, columns):
    def cell(column):
        index = columns.get(column)
        return row[index] if index is not None and index < len(row) else None

    if is_ticked(cell("sent")):
        return StatusPhase.SENT_TO_CLIENT
    if is_ticked(cell("reviewed")):
        return StatusPhase.REVIEWED
    if is_ticked(cell("submitted")):
        return StatusPhase.SUBMITTED
    return STATUS_TEXT.get(header_key(cell("status")), StatusPhase.PENDING)


def import_tracker_workbook(uploaded_file, quarter, remove_missing=False, year=None):
    year = year or date.today().year
    try:
        workbook = load_workbook(uploaded_file, data_only=True, read_only=True)
    except Exception as exc:
        raise TrackerUploadError("The file could not be read. Upload an Excel workbook (.xlsx).") from exc

    worksheet, header_row, columns = find_header(workbook)
    result = TrackerUploadResult(sheet_name=worksheet.title)

    existing = {name_key(report.client_name): report for report in ClientReport.objects.filter(quarter=quarter)}
    seen = set()

    for row in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_name = row[columns["client"]] if columns["client"] < len(row) else None
        if not isinstance(raw_name, str) or not raw_name.strip():
            continue
        client_name = " ".join(raw_name.split())
        if client_name.upper().startswith(("STATUS:", "TOTAL")):
            continue

        key = name_key(client_name)
        if key in CLIENT_NAME_ALIASES:
            client_name = CLIENT_NAME_ALIASES[key]
            key = name_key(client_name)
        if key in seen:
            result.skipped.append(f"{client_name} (listed more than once)")
            continue
        seen.add(key)

        def text(column):
            index = columns.get(column)
            value = row[index] if index is not None and index < len(row) else None
            return " ".join(str(value).split()) if value not in (None, "") else ""

        values = {
            "crm_name": text("crm"),
            "operations_assignee": text("operations"),
            "status_phase": status_from_row(row, columns),
        }
        if "due_date" in columns:
            values["due_date"] = parse_due_date(row[columns["due_date"]], year)
        if "notes" in columns:
            values["notes"] = text("notes")[:255]

        report = existing.get(key)
        if report is None:
            report = ClientReport(quarter=quarter, client_name=client_name)
            result.created.append(client_name)
        else:
            result.updated.append(report.client_name)
        for name, value in values.items():
            setattr(report, name, value)
        report.save()

    if not seen:
        raise TrackerUploadError(f"No client rows were found under the headers on sheet '{worksheet.title}'.")

    if remove_missing:
        for key, report in existing.items():
            if key not in seen:
                result.removed.append(report.client_name)
                report.delete()

    return result
